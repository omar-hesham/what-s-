"""
Local document extraction pipeline.
Supports PDF, DOCX, XLSX, TXT, and CSV with bounded limits.
Explicitly marks legacy .doc and .xls as unsupported without fake parsing.
Never fabricates document text or embeds error strings into extracted content.
Preserves document provenance and indexes verified text in derived_fts.
"""

from datetime import datetime
from pathlib import Path
from typing import Dict, Any
from sqlalchemy.orm import Session
from pypdf import PdfReader
from docx import Document as DocxDocument
import openpyxl

from owi.core.logging import logger
from owi.db.models import MediaAsset, DocumentRecord
from owi.db.migrations import sync_derived_fts, remove_derived_fts

# Resource limits to prevent OOM
MAX_DOC_BYTES = 5 * 1024 * 1024  # 5 MB for plain text/csv
MAX_PDF_PAGES = 100
MAX_DOCX_PARAGRAPHS = 2000
MAX_XLSX_SHEETS = 10
MAX_XLSX_ROWS_PER_SHEET = 1000

class DocumentProcessor:
    """Extracts bounded text and metadata from PDF, Word, Excel, and text documents."""

    @classmethod
    def process_document(cls, media_asset_id: int, db: Session) -> Dict[str, Any]:
        asset = db.query(MediaAsset).filter(MediaAsset.id == media_asset_id).first()
        if not asset:
            raise ValueError(f"MediaAsset #{media_asset_id} not found.")

        file_path = Path(asset.file_path)
        if not file_path.exists():
            asset.processing_status = "failed"
            asset.processing_error = "file_not_found"
            asset.processing_attempts = (asset.processing_attempts or 0) + 1
            asset.processed_at = datetime.utcnow()
            db.commit()
            raise FileNotFoundError(f"Document file does not exist: {file_path}")

        asset.processing_attempts = (asset.processing_attempts or 0) + 1
        asset.processing_status = "processing"
        db.commit()

        suffix = file_path.suffix.lower()

        # 1. Explicitly check for legacy unsupported formats
        if suffix in (".doc", ".xls"):
            logger.info(f"Legacy format '{suffix}' is not supported locally: {file_path.name}")
            asset.processing_status = "unsupported"
            asset.processing_error = f"legacy_binary_format_unsupported ({suffix})"
            asset.processing_method = "none"
            asset.processed_at = datetime.utcnow()
            remove_derived_fts(db.connection(), asset.id)
            db.commit()
            return {
                "media_asset_id": asset.id,
                "status": "unsupported",
                "error": f"Legacy {suffix} files are unsupported locally. Please convert to modern {suffix}x format.",
                "extracted_text": "",
                "char_count": 0
            }

        extracted_text = ""
        page_count = 1
        meta: Dict[str, Any] = {}

        try:
            if suffix == ".pdf":
                reader = PdfReader(file_path)
                total_pages = len(reader.pages)
                page_count = min(total_pages, MAX_PDF_PAGES)
                pages_text = []
                for p_idx in range(page_count):
                    page = reader.pages[p_idx]
                    text = (page.extract_text() or "").strip()
                    if text:
                        pages_text.append(f"--- Page {p_idx + 1} ---\n{text}")
                extracted_text = "\n\n".join(pages_text)
                if reader.metadata:
                    meta = {k: str(v) for k, v in reader.metadata.items() if v}
                meta["total_pages_detected"] = total_pages

            elif suffix == ".docx":
                doc = DocxDocument(file_path)
                paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()][:MAX_DOCX_PARAGRAPHS]
                # Extract bounded tables
                table_texts = []
                for table in doc.tables[:50]:
                    for row in table.rows[:200]:
                        row_vals = [cell.text.strip() for cell in row.cells]
                        if any(row_vals):
                            table_texts.append(" | ".join(row_vals))

                parts = []
                if paragraphs:
                    parts.append("\n".join(paragraphs))
                if table_texts:
                    parts.append("--- Tables ---\n" + "\n".join(table_texts))
                extracted_text = "\n\n".join(parts)
                page_count = max(1, len(paragraphs) // 10)

            elif suffix == ".xlsx":
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                sheets_text = []
                sheet_names = wb.sheetnames[:MAX_XLSX_SHEETS]
                for sheetname in sheet_names:
                    sheet = wb[sheetname]
                    sheet_lines = [f"=== Sheet: {sheetname} ==="]
                    row_count = 0
                    for row in sheet.iter_rows(values_only=True):
                        if row_count >= MAX_XLSX_ROWS_PER_SHEET:
                            sheet_lines.append(f"... (truncated after {MAX_XLSX_ROWS_PER_SHEET} rows)")
                            break
                        if any(cell is not None for cell in row):
                            row_str = " | ".join(str(cell) if cell is not None else "" for cell in row[:50])
                            if row_str.strip():
                                sheet_lines.append(row_str)
                                row_count += 1
                    sheets_text.append("\n".join(sheet_lines))
                wb.close()
                extracted_text = "\n\n".join(sheets_text)
                meta["sheet_names"] = sheet_names
                page_count = len(sheet_names)

            elif suffix in (".csv", ".txt"):
                # Bounded read to avoid memory exhaustion
                f_size = file_path.stat().st_size
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    extracted_text = f.read(MAX_DOC_BYTES)
                if f_size > MAX_DOC_BYTES:
                    meta["truncated"] = True
                    meta["original_size"] = f_size
                page_count = max(1, extracted_text.count("\n") // 40)

            else:
                asset.processing_status = "unsupported"
                asset.processing_error = f"unsupported_document_extension ({suffix})"
                asset.processing_method = "none"
                asset.processed_at = datetime.utcnow()
                db.commit()
                return {
                    "media_asset_id": asset.id,
                    "status": "unsupported",
                    "error": f"Unsupported document extension: {suffix}",
                    "extracted_text": "",
                    "char_count": 0
                }

        except Exception as e:
            logger.error(f"Error extracting document text from {file_path.name}: {e}")
            asset.processing_status = "failed"
            asset.processing_error = f"extraction_error: {str(e)[:150]}"
            asset.processing_method = f"local-{suffix.lstrip('.')}"
            asset.processed_at = datetime.utcnow()
            db.commit()
            return {
                "media_asset_id": asset.id,
                "status": "failed",
                "error": str(e),
                "extracted_text": "",
                "char_count": 0
            }

        # Successful extraction
        asset.processing_status = "completed"
        asset.processing_method = f"local-{suffix.lstrip('.')}"
        asset.processing_error = None
        asset.processed_at = datetime.utcnow()

        # Idempotently update DocumentRecord
        doc_record = db.query(DocumentRecord).filter(DocumentRecord.media_asset_id == asset.id).first()
        if not doc_record:
            doc_record = DocumentRecord(
                media_asset_id=asset.id,
                conversation_id=asset.conversation_id,
                title=file_path.stem,
                doc_type=suffix.lstrip("."),
                page_count=page_count,
                extracted_text=extracted_text,
                metadata_json=meta
            )
            db.add(doc_record)
        else:
            doc_record.title = file_path.stem
            doc_record.doc_type = suffix.lstrip(".")
            doc_record.page_count = page_count
            doc_record.extracted_text = extracted_text
            doc_record.metadata_json = meta

        db.flush()

        # Sync verified non-empty text to derived_fts
        if extracted_text.strip():
            sync_derived_fts(
                db.connection(),
                media_asset_id=asset.id,
                message_id=asset.message_id,
                conversation_id=asset.conversation_id,
                file_name=asset.file_name,
                source_type="document",
                content=extracted_text
            )
        else:
            remove_derived_fts(db.connection(), asset.id)

        db.commit()

        return {
            "media_asset_id": asset.id,
            "document_id": doc_record.id,
            "status": "completed",
            "title": doc_record.title,
            "doc_type": doc_record.doc_type,
            "page_count": doc_record.page_count,
            "char_count": len(extracted_text)
        }
